#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "data" / "standardized" / "manual_alternative_observations.csv"
REPORT = PROJECT_ROOT / "data" / "processed" / "tracking_workbook_import_report.json"

FIELDS = [
    "series_id", "entity_name", "ticker", "metric_name", "observation_end",
    "available_at", "metric_value", "unit", "frequency", "source_name",
    "source_url", "source_tier", "pit_status", "workbook_sheet", "workbook_cell",
    "research_use", "notes",
]


def iso_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        text = value.strip()
        match = re.search(r"(20\d{2})\s*[- ]?\s*(\d{1,2})\s*月?", text)
        if match:
            return date(int(match.group(1)), int(match.group(2)), 1).isoformat()
        match = re.search(r"(20\d{2}).*?Q([1-4])", text, re.I)
        if match:
            year, quarter = int(match.group(1)), int(match.group(2))
            month = quarter * 3
            day = 31 if month in (3, 12) else 30
            return date(year, month, day).isoformat()
        match = re.search(r"(20\d{2})年初", text)
        if match:
            return date(int(match.group(1)), 1, 1).isoformat()
        match = re.search(r"(20\d{2})年全年", text)
        if match:
            return date(int(match.group(1)), 12, 31).isoformat()
    return None


def quarter_end(value) -> str | None:
    match = re.fullmatch(r"(20\d{2})Q([1-4])", str(value).strip(), re.I)
    if not match:
        return None
    year, quarter = int(match.group(1)), int(match.group(2))
    month = quarter * 3
    return date(year, month, 31 if month in (3, 12) else 30).isoformat()


def add(rows: list[dict], *, series_id: str, entity: str, metric: str,
        observation_end: str | None, value, available_at: str, unit: str,
        frequency: str, source_name: str, source_url: str, source_tier: str,
        sheet: str, cell: str, ticker: str = "", notes: str = "",
        pit_status: str = "snapshot_available_only",
        research_use: str = "descriptive_only") -> None:
    if observation_end is None or not isinstance(value, (int, float)):
        return
    rows.append({
        "series_id": series_id,
        "entity_name": entity,
        "ticker": ticker,
        "metric_name": metric,
        "observation_end": observation_end,
        "available_at": available_at,
        "metric_value": value,
        "unit": unit,
        "frequency": frequency,
        "source_name": source_name,
        "source_url": source_url,
        "source_tier": source_tier,
        "pit_status": pit_status,
        "workbook_sheet": sheet,
        "workbook_cell": cell,
        "research_use": research_use,
        "notes": notes,
    })


def main() -> int:
    parser = argparse.ArgumentParser(description="Import curated AI tracking workbook observations.")
    parser.add_argument("--input", required=True, type=Path)
    args = parser.parse_args()
    source = args.input.expanduser().resolve()
    workbook = openpyxl.load_workbook(source, read_only=False, data_only=True)
    snapshot_at = datetime.fromtimestamp(source.stat().st_mtime).astimezone().isoformat()
    rows: list[dict] = []

    sheet = workbook["算力租赁价格"]
    gpu_columns = {3: ("H100", "gpu_rental_h100_usd_hour"), 4: ("A100", "gpu_rental_a100_usd_hour"), 5: ("B200", "gpu_rental_b200_usd_hour")}
    for row in range(2, sheet.max_row + 1):
        observed = iso_date(sheet.cell(row, 2).value)
        for column, (entity, series_id) in gpu_columns.items():
            add(rows, series_id=series_id, entity=entity, metric="gpu_rental_price",
                observation_end=observed, value=sheet.cell(row, column).value,
                available_at=snapshot_at, unit="USD_per_hour", frequency="daily",
                source_name="source_not_recorded_in_workbook", source_url="", source_tier="U",
                sheet=sheet.title, cell=sheet.cell(row, column).coordinate,
                notes="Useful market proxy; vendor and historical publication timestamps require validation.")

    sheet = workbook["存储价格"]
    memory_blocks = [
        (2, {3: "DDR5_16Gb_2Gx8", 4: "DDR5_16Gb_1Gx16", 5: "DDR4_8Gb_1Gx8", 6: "DDR4_4Gb_512Mx8"}),
        (8, {9: "QLC_Flash_1Tb", 10: "TLC_Flash_1Tb", 11: "TLC_Flash_512Gb", 12: "TLC_Flash_256Gb"}),
    ]
    for date_column, series in memory_blocks:
        for row in range(2, sheet.max_row + 1):
            observed = iso_date(sheet.cell(row, date_column).value)
            for column, entity in series.items():
                add(rows, series_id=f"memory_spot_{entity.lower()}", entity=entity,
                    metric="memory_spot_price", observation_end=observed,
                    value=sheet.cell(row, column).value, available_at=snapshot_at,
                    unit="workbook_native", frequency="daily",
                    source_name="source_not_recorded_in_workbook", source_url="", source_tier="U",
                    sheet=sheet.title, cell=sheet.cell(row, column).coordinate,
                    notes="Unit and vendor are not documented; use percentage changes only until validated.")

    sheet = workbook["北美四大云 CSP CAPEX"]
    capex_columns = {
        21: ("Microsoft", "MSFT", "https://www.microsoft.com/en-us/investor/sec-filings"),
        22: ("Alphabet", "GOOGL", "https://abc.xyz/investor/default.aspx"),
        23: ("Amazon", "AMZN", "https://ir.aboutamazon.com/quarterly-results/default.aspx"),
        24: ("Meta", "META", "https://investor.atmeta.com/home/default.aspx"),
        25: ("Top 4 CSP", "", ""),
    }
    for row in range(2, sheet.max_row + 1):
        observed = quarter_end(sheet.cell(row, 20).value)
        for column, (entity, ticker, url) in capex_columns.items():
            add(rows, series_id=f"csp_capex_{ticker.lower() or 'top4'}", entity=entity,
                ticker=ticker, metric="cash_capex", observation_end=observed,
                value=sheet.cell(row, column).value, available_at=snapshot_at,
                unit="USD_bn", frequency="quarterly", source_name="company_investor_relations_compilation",
                source_url=url, source_tier="B", sheet=sheet.title,
                cell=sheet.cell(row, column).coordinate,
                notes="Official IR landing page is recorded, but exact filing publication timestamps were not retained.")

    sheet = workbook["AI MAU"]
    activity_blocks = [
        (range(2, 9), 3, "ChatGPT", "weekly_active_users", "million_users", "technologychecker.io / OpenAI statements", "https://technologychecker.io/blog/chatgpt-statistics", "B"),
        (range(24, 31), 3, "Gemini", "monthly_active_users", "million_users", "ClickVision / Google statements", "https://click-vision.com/google-gemini-statistics", "B"),
    ]
    for row_range, column, entity, metric, unit, source_name, url, tier in activity_blocks:
        for row in row_range:
            verified_latest_chatgpt = entity == "ChatGPT" and row == 8
            add(rows, series_id=f"product_activity_{entity.lower()}", entity=entity, metric=metric,
                observation_end=iso_date(sheet.cell(row, 2).value), value=sheet.cell(row, column).value,
                available_at="2026-02-27T00:00:00+00:00" if verified_latest_chatgpt else snapshot_at,
                unit=unit, frequency="irregular",
                source_name="OpenAI" if verified_latest_chatgpt else source_name,
                source_url="https://openai.com/index/scaling-ai-for-everyone/" if verified_latest_chatgpt else url,
                source_tier="A" if verified_latest_chatgpt else tier, sheet=sheet.title,
                cell=sheet.cell(row, column).coordinate,
                notes="Official OpenAI disclosure cross-verified; usable prospectively from publication date."
                if verified_latest_chatgpt else "Historical source-release timestamps are not retained; descriptive product-adoption proxy only.",
                pit_status="source_publication_verified" if verified_latest_chatgpt else "snapshot_available_only",
                research_use="descriptive_and_forward_pit" if verified_latest_chatgpt else "descriptive_only")
    china_products = {3: "DeepSeek", 4: "Doubao", 5: "Yuanbao", 6: "Qwen", 7: "Kimi"}
    for row in range(47, 53):
        observed = iso_date(sheet.cell(row, 2).value)
        for column, entity in china_products.items():
            add(rows, series_id=f"product_activity_{entity.lower()}", entity=entity,
                metric="monthly_active_users", observation_end=observed,
                value=sheet.cell(row, column).value, available_at=snapshot_at,
                unit="million_users", frequency="irregular", source_name="QuestMobile",
                source_url="https://www.questmobile.com.cn/research/", source_tier="B",
                sheet=sheet.title, cell=sheet.cell(row, column).coordinate,
                notes="Vendor-estimated MAU; methodology and historical availability require separate snapshots.")

    sheet = workbook["Tokens消耗量"]
    for row in range(24, 27):
        add(rows, series_id="google_monthly_token_usage", entity="Google", ticker="GOOGL",
            metric="token_usage", observation_end=iso_date(sheet.cell(row, 2).value),
            value=sheet.cell(row, 3).value, available_at=snapshot_at, unit="trillion_tokens_per_month",
            frequency="irregular", source_name="NewPages", source_url="https://www.newpages.net/latestnews/nid/185480/",
            source_tier="C", sheet=sheet.title, cell=sheet.cell(row, 3).coordinate,
            notes="Secondary-source company usage estimate; not suitable for PIT backtests.")
    for row in range(52, 60):
        add(rows, series_id="doubao_daily_token_usage", entity="Doubao",
            metric="token_usage", observation_end=iso_date(sheet.cell(row, 2).value),
            value=sheet.cell(row, 3).value, available_at=snapshot_at, unit="trillion_tokens_per_day",
            frequency="irregular", source_name="Volcengine official WeChat",
            source_url="https://mp.weixin.qq.com/s/wypvdbPmVR8XH4y4G8mS7Q", source_tier="B",
            sheet=sheet.title, cell=sheet.cell(row, 3).coordinate,
            notes="Company-reported operating metric; exact announcement URL may vary by observation.")

    rows.sort(key=lambda item: (item["series_id"], item["observation_end"], item["workbook_cell"]))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    by_metric: dict[str, int] = {}
    by_tier: dict[str, int] = {}
    for item in rows:
        by_metric[item["metric_name"]] = by_metric.get(item["metric_name"], 0) + 1
        by_tier[item["source_tier"]] = by_tier.get(item["source_tier"], 0) + 1
    report = {
        "source_file": source.name,
        "source_modified_at": snapshot_at,
        "rows_written": len(rows),
        "series_count": len({item["series_id"] for item in rows}),
        "rows_by_metric": by_metric,
        "rows_by_source_tier": by_tier,
        "backtest_eligible_rows": sum(item["pit_status"] == "source_publication_verified" for item in rows),
        "policy": "Imported values are available no earlier than the workbook snapshot timestamp unless an original publication date is independently verified. Historical vintages must be collected before broad backtest use.",
        "excluded_for_now": ["model ARR", "token price", "CDS", "financing narratives", "NVDA consensus expectations"],
        "exclusion_reason": "Mixed or incomplete provenance, ambiguous units, proprietary-source dependency, or missing point-in-time availability.",
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"rows={len(rows)} series={report['series_count']} output={OUTPUT.relative_to(PROJECT_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
